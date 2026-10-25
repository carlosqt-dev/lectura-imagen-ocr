"""Escribe un ResultadoProcesado en la plantilla .xlsx/.xlsm de F-LA-08-01."""
from __future__ import annotations

import logging
from pathlib import Path

import openpyxl

from models import ResultadoProcesado

from .mapeo_celdas import MAPEO_CELDAS, NOMBRE_HOJA, VALIDACION_ENCABEZADOS

logger = logging.getLogger(__name__)


class ExcelWriteError(Exception):
    """No fue posible abrir o guardar la plantilla de Excel."""


def _a_numero(valor: str) -> float | str:
    """El OCR entrega los números como texto, a veces con coma decimal."""
    try:
        return float(valor.replace(",", "."))
    except ValueError:
        return valor


def escribir_resultado(
    resultado: ResultadoProcesado, ruta_plantilla: Path, ruta_salida: Path
) -> Path:
    """Vuelca los 7 campos transcritos en la plantilla y guarda una copia en ruta_salida.

    Los campos cuyo valor es None (pendientes) se dejan vacíos y sólo generan
    un warning; nunca se detiene el guardado del resto del archivo. El resto
    de las celdas de la plantilla (las que diligencia el operario a mano) no
    se tocan.
    """
    try:
        libro = openpyxl.load_workbook(ruta_plantilla, keep_vba=True)
    except Exception as exc:
        raise ExcelWriteError(f"No se pudo abrir la plantilla: {ruta_plantilla}") from exc

    if NOMBRE_HOJA not in libro.sheetnames:
        raise ExcelWriteError(f"La plantilla no tiene la hoja '{NOMBRE_HOJA}'.")
    hoja = libro[NOMBRE_HOJA]

    # Los encabezados de la fila 17 son fórmulas (VLOOKUP); para validar su
    # texto resultante hace falta el valor calculado, no la fórmula en sí,
    # así que se inspecciona con un libro aparte abierto en modo data_only.
    libro_valores = openpyxl.load_workbook(ruta_plantilla, data_only=True)
    hoja_valores = libro_valores[NOMBRE_HOJA]
    for clave, (celda_encabezado, texto_esperado) in VALIDACION_ENCABEZADOS.items():
        actual = hoja_valores[celda_encabezado].value
        if actual != texto_esperado:
            logger.warning(
                "La plantilla cambió: %s!%s dice %r (se esperaba %r). "
                "Revise si el tipo de probeta/material (K14/K15) cambió antes de confiar en '%s'.",
                NOMBRE_HOJA, celda_encabezado, actual, texto_esperado, clave,
            )

    datos = resultado.resultado.como_dict()
    for clave, celda in MAPEO_CELDAS.items():
        valor = datos.get(clave)
        if valor is None:
            logger.warning("Campo pendiente, celda vacía: %s (%s!%s)", clave, NOMBRE_HOJA, celda)
            continue
        hoja[celda] = _a_numero(valor)

    try:
        ruta_salida.parent.mkdir(parents=True, exist_ok=True)
        libro.save(ruta_salida)
    except Exception as exc:
        raise ExcelWriteError(f"No se pudo guardar el archivo: {ruta_salida}") from exc

    return ruta_salida
